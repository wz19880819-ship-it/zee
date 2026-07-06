"""Generate the argument paper as an APA-formatted Word document."""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION

TITLE = "Joining the Military After High School: A Good Choice for Young People"

BODY = [
    # (is_heading, text)
    (False, "Young people in America today face many serious problems. They deal with "
            "school stress, family problems, drugs, violence, and trouble finding good jobs "
            "(TopTenz, n.d.). Many teenagers finish high school without a clear plan for their "
            "future. One possible answer to this problem is very controversial: military "
            "recruiters visit high schools and encourage young people to join the military. "
            "Critics say that recruiters take advantage of poor students who have few choices. "
            "Supporters say that the military gives young people money for college, job "
            "training, and discipline. Although some people believe that military recruiting "
            "uses young people unfairly, I argue that joining the military, such as the Army "
            "National Guard, is a positive choice for many young Americans because it pays "
            "for education, teaches real job skills, and builds discipline and maturity."),
    (False, "First, it is important to understand the controversy. Under federal law, public "
            "high schools that receive government money must give military recruiters the "
            "same access to students that they give to colleges and employers. Recruiters can "
            "also receive student contact information unless parents say no. Critics, such as "
            "the American Civil Liberties Union (2008), argue that this system targets "
            "students from low-income families and minority communities, and that some "
            "recruiters make promises they cannot keep. On the other side, military leaders "
            "say that recruiting is necessary because the United States has an all-volunteer "
            "force, and young people deserve to hear about military service as one option "
            "among many (Segal & Segal, 2004). The question is simple but serious: does "
            "military recruiting help young people or hurt them?"),
    (False, "My first reason for supporting military service is education money. College in "
            "America is very expensive, and many families cannot pay for it. The military "
            "offers the GI Bill and tuition assistance, which can pay for most or all of a "
            "college degree. Research shows that many young people join the military exactly "
            "for this reason. Kleykamp (2006) studied high school students and found that "
            "students who wanted to go to college, but did not have enough money, often chose "
            "the military as a bridge to higher education. In the National Guard, this benefit "
            "is even more flexible. Guard soldiers serve part time, usually one weekend each "
            "month, so they can go to college and serve at the same time. For a young person "
            "with no savings, this is a real path to a degree without heavy student loans."),
    (False, "My second reason is job training. The military is not only about combat. Most "
            "military jobs are support jobs, such as logistics, medical care, communications, "
            "and engineering. I know this from my own experience. I serve as a Staff Sergeant "
            "in the Army National Guard, working in J4 logistics. In my job, I learned supply "
            "chain management, inventory control, transportation planning, and how to lead a "
            "team under pressure. These are the same skills that big companies like Amazon, "
            "FedEx, and Walmart look for every day. Segal and Segal (2004) explain that the "
            "modern military is a huge training organization that prepares young people for "
            "technical and management careers. A nineteen-year-old who spends a few years in "
            "military logistics can leave with skills, certifications, and leadership "
            "experience that most college students do not have."),
    (False, "My third reason is discipline and personal growth. Many young people feel lost "
            "after high school. The military gives structure: wake up early, stay physically "
            "fit, be on time, respect others, and finish the mission. Research supports this "
            "idea. Elder (1986) followed men over many years and found that military service "
            "was often a positive turning point, especially for young men from disadvantaged "
            "backgrounds. Service gave them time to grow up, new social networks, and a "
            "second chance to build a better life. I have seen this with my own soldiers. "
            "Young people arrive shy and disorganized, and after training they stand taller, "
            "speak clearly, and take responsibility for their team."),
    (False, "However, honest research means listening to the other side. Critics make strong "
            "points. The American Civil Liberties Union (2008) reported that some recruiters "
            "used aggressive methods with students under eighteen, and that recruiting "
            "stations are placed more often in poor neighborhoods. Critics also remind us "
            "that military service has real risks, including deployment, injury, and mental "
            "health problems like PTSD. These concerns are serious, and I do not ignore "
            "them. But I believe the answer is honest recruiting, not less opportunity. "
            "Recruiters should tell the full truth about risks and benefits, and parents "
            "should stay involved in the decision. Also, the National Guard model lowers "
            "many of these risks, because most service is part time and close to home, "
            "helping local communities during storms, floods, and other emergencies. Taking "
            "away military options from poor students does not give them college money; it "
            "only removes one of the few doors that is open to them."),
    (False, "In conclusion, the controversy about military recruiting in high schools is "
            "really a question about what is best for our youth. Critics are right that "
            "recruiting must be honest and fair. But the evidence, and my own service, show "
            "that the military offers young people money for education, valuable job skills, "
            "and the discipline to succeed in life. For many young Americans who feel stuck "
            "after high school, raising their right hand is not a trap. It is a beginning."),
]

REFERENCES = [
    "American Civil Liberties Union. (2008). Soldiers of misfortune: Abusive U.S. military "
    "recruitment and failure to protect child soldiers. https://www.aclu.org/report/"
    "soldiers-misfortune-abusive-us-military-recruitment-and-failure-protect-child-soldiers",

    "Elder, G. H., Jr. (1986). Military times and turning points in men's lives. "
    "Developmental Psychology, 22(2), 233–245. https://doi.org/10.1037/0012-1649.22.2.233",

    "Kleykamp, M. A. (2006). College, jobs, or the military? Enlistment during a time of "
    "war. Social Science Quarterly, 87(2), 272–290. https://doi.org/10.1111/"
    "j.1540-6237.2006.00380.x",

    "Segal, D. R., & Segal, M. W. (2004). America's military population. Population "
    "Bulletin, 59(4), 1–40.",

    "TopTenz. (n.d.). Top 10 issues facing our youth today. "
    "https://www.toptenz.net/top-10-issues-facing-our-youth-today.php",
]

doc = Document()

# Global style: Times New Roman 12, double spacing
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)
style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
style.paragraph_format.space_after = Pt(0)
style.paragraph_format.space_before = Pt(0)

# 1-inch margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def para(text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, indent=None):
    p = doc.add_paragraph()
    p.alignment = align
    if indent is not None:
        p.paragraph_format.first_line_indent = Inches(indent)
    run = p.add_run(text)
    run.bold = bold
    return p

# ---------- Cover page (APA student style) ----------
for _ in range(7):
    doc.add_paragraph()
para(TITLE, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
doc.add_paragraph()
for line in ["Ze Wang", "Berkeley College", "The Confident Writer 1264",
             "Professor [Instructor Name]", "July 6, 2026"]:
    para(line, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ---------- Body ----------
para(TITLE, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
for _, text in BODY:
    para(text, indent=0.5)
doc.add_page_break()

# ---------- References ----------
para("References", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
for ref in REFERENCES:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)  # hanging indent
    p.add_run(ref)

out = "/home/user/zee/Ze_Wang_Argument_Paper.docx"
doc.save(out)

words = sum(len(t.split()) for _, t in BODY)
print(f"Saved {out}  |  body word count: {words}")
